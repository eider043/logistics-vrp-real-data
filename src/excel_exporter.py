"""
Exportador Plan Operativo Excel — datos reales
Autor: Eider
"""

import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
import os

os.makedirs("../outputs", exist_ok=True)

PASTEL = {
    "azul":   "AED6F1", "verde":  "A9DFBF",
    "amarillo":"F9E79F", "morado": "D7BDE2",
    "salmon": "F1948A", "gris":   "F2F3F4",
    "header": "5D8AA8", "blanco": "FFFFFF",
}

def border():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)


def exportar_plan_operativo(viajes, kpis, scenario, nombre_escenario="Real"):
    fecha = datetime.now().strftime("%Y%m%d_%H%M")
    path  = f"../outputs/plan_operativo_{nombre_escenario}_{fecha}.xlsx"
    wb    = openpyxl.Workbook()

    # ── Hoja 1: Resumen ───────────────────────────────────────────────
    ws = wb.active
    ws.title = "Resumen Ejecutivo"
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:I1")
    c = ws["A1"]
    c.value = f"PLAN OPERATIVO DE DISTRIBUCION — {nombre_escenario.upper()} | Datos Reales Fleet Operations"
    c.font  = Font(bold=True, size=13, color="FFFFFF")
    c.fill  = PatternFill("solid", fgColor=PASTEL["header"])
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:I2")
    ws["A2"].value = f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')} | Ordenes procesadas: {sum(len(v['ordenes']) for v in viajes)}"
    ws["A2"].alignment = Alignment(horizontal="center")
    ws["A2"].font = Font(size=9, color="666666")

    kpi_data = [
        ("Total Viajes",        kpis["total_viajes"],                  PASTEL["azul"]),
        ("Costo Total Fletes",  f"${kpis['costo_total']:,.0f}",       PASTEL["salmon"]),
        ("Revenue Total",       f"${kpis['revenue_total']:,.0f}",     PASTEL["verde"]),
        ("Margen Bruto",        f"${kpis['margen_total']:,.0f}",      PASTEL["morado"]),
        ("Peso Movilizado",     f"{kpis['peso_total_kg']:,.0f} kg",   PASTEL["amarillo"]),
        ("Distancia Total",     f"{kpis['dist_total_km']:,.0f} km",   PASTEL["azul"]),
        ("Ocup. Peso Prom.",    f"{kpis['ocup_peso_prom']}%",         PASTEL["verde"]),
        ("Costo/kg Prom.",      f"${kpis['costo_por_kg_prom']:.4f}",  PASTEL["salmon"]),
        ("Rentabilidad",        f"{kpis['margen_total']/max(kpis['revenue_total'],1)*100:.1f}%", PASTEL["morado"]),
    ]

    for col, (titulo, valor, color) in enumerate(kpi_data, 1):
        ws.column_dimensions[get_column_letter(col)].width = 18
        c = ws.cell(5, col, titulo)
        c.font = Font(bold=True, size=9, color="333333")
        c.fill = PatternFill("solid", fgColor=color)
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.border = border()
        ws.row_dimensions[5].height = 25
        c2 = ws.cell(6, col, str(valor))
        c2.font = Font(bold=True, size=13)
        c2.fill = PatternFill("solid", fgColor="FDFEFE")
        c2.alignment = Alignment(horizontal="center", vertical="center")
        c2.border = border()
        ws.row_dimensions[6].height = 32

    # ── Hoja 2: Plan de Viajes ────────────────────────────────────────
    ws2 = wb.create_sheet("Plan de Viajes")
    ws2.sheet_view.showGridLines = False

    headers = [
        "Viaje ID", "Hub Origen", "Ciudad Hub",
        "Cliente Principal", "Ciudad Destino",
        "N. Paradas", "Distancia (km)",
        "Peso (kg)", "Volumen (m3)",
        "Ocup. Peso %", "Ocup. Vol %",
        "Vehiculo", "Cap. Peso (kg)", "Cap. Vol (m3)",
        "Transportadora", "Costo Flete", "Revenue", "Margen",
        "Costo/kg", "Ordenes"
    ]

    for col, h in enumerate(headers, 1):
        c = ws2.cell(1, col, h)
        c.font = Font(bold=True, size=9, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=PASTEL["header"])
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.border = border()
    ws2.row_dimensions[1].height = 28

    widths = [10,18,14,20,14,9,13,10,10,11,9,18,13,11,22,14,14,14,10,25]
    for i, w in enumerate(widths, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    colores_alt = [PASTEL["blanco"], PASTEL["gris"]]
    for r, v in enumerate(viajes, 2):
        color = colores_alt[r % 2]
        row = [
            v["viaje_id"], v["cedi_nombre"], v["cedi_ciudad"],
            v["cliente_nombre"], v["cliente_ciudad"],
            v["n_paradas"], v["distancia_km"],
            v["peso_kg"], v["vol_m3"],
            v["ocup_peso_pct"], v["ocup_vol_pct"],
            v["vehiculo"], v["cap_peso_kg"], v["cap_vol_m3"],
            v["transportadora"], v["costo_total"], v["revenue_viaje"], v["margen_viaje"],
            v["costo_por_kg"], ", ".join(v["ordenes"][:3]) + ("..." if len(v["ordenes"]) > 3 else "")
        ]
        for col, val in enumerate(row, 1):
            c = ws2.cell(r, col, val)
            c.fill = PatternFill("solid", fgColor=color)
            c.border = border()
            c.font = Font(size=9)
            c.alignment = Alignment(vertical="center")
            if col in [16, 17, 18]:
                c.number_format = "#,##0"
            if col in [10, 11] and isinstance(val, float):
                if val < 30:   c.fill = PatternFill("solid", fgColor=PASTEL["salmon"])
                elif val > 70: c.fill = PatternFill("solid", fgColor=PASTEL["verde"])
            if col == 18 and isinstance(val, float):
                c.font = Font(size=9, color="1E8449" if val > 0 else "C0392B", bold=True)
        ws2.row_dimensions[r].height = 20

    # Totales
    tr = len(viajes) + 2
    for col, val in enumerate([
        "TOTAL", "", "", "", "", "", round(kpis["dist_total_km"],1),
        round(kpis["peso_total_kg"],1), round(kpis["vol_total_m3"],2),
        kpis["ocup_peso_prom"], kpis["ocup_vol_prom"],
        "", "", "", "",
        round(kpis["costo_total"],0), round(kpis["revenue_total"],0), round(kpis["margen_total"],0),
        "", ""
    ], 1):
        c = ws2.cell(tr, col, val)
        c.font = Font(bold=True, size=9)
        c.fill = PatternFill("solid", fgColor=PASTEL["azul"])
        c.border = border()

    # ── Hoja 3: Por Hub ───────────────────────────────────────────────
    ws3 = wb.create_sheet("Por Hub")
    ws3.sheet_view.showGridLines = False
    df = pd.DataFrame(viajes)
    resumen = df.groupby(["cedi_id","cedi_nombre","cedi_ciudad"]).agg(
        viajes=("viaje_id","count"),
        costo=("costo_total","sum"),
        revenue=("revenue_viaje","sum"),
        margen=("margen_viaje","sum"),
        peso=("peso_kg","sum"),
        dist=("distancia_km","sum"),
        ocup_peso=("ocup_peso_pct","mean"),
    ).reset_index()

    hs = ["Hub ID","Nombre Hub","Ciudad","Viajes","Costo Fletes","Revenue","Margen","Peso (kg)","Dist. (km)","Ocup. Peso %"]
    for col, h in enumerate(hs, 1):
        c = ws3.cell(1, col, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=PASTEL["header"])
        c.border = border()

    for r, row in resumen.iterrows():
        vals = [row.cedi_id, row.cedi_nombre, row.cedi_ciudad,
                row.viajes, round(row.costo,0), round(row.revenue,0),
                round(row.margen,0), round(row.peso,1), round(row.dist,1), round(row.ocup_peso,1)]
        for col, val in enumerate(vals, 1):
            c = ws3.cell(r+2, col, val)
            c.border = border()
            c.fill = PatternFill("solid", fgColor=colores_alt[r%2])

    for i, w in enumerate([10,18,14,8,15,15,14,12,11,12], 1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    wb.save(path)
    print(f"Plan exportado: {path}")
    return path


if __name__ == "__main__":
    import sys
    sys.path.insert(0, ".")
    from data_loader import load_freight_data, build_scenario
    from optimizer import optimizar_escenario
    df = load_freight_data()
    scenario = build_scenario(df, 50)
    viajes, kpis = optimizar_escenario(scenario)
    exportar_plan_operativo(viajes, kpis, scenario)